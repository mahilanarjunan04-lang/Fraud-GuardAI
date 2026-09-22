import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    NAVY_BG = RGBColor(7, 18, 36)        # #071224
    PANEL_BG = RGBColor(14, 30, 56)      # #0E1E38
    EMERALD = RGBColor(16, 185, 129)     # #10B981
    ROYAL_BLUE = RGBColor(29, 78, 216)   # #1D4ED8
    CYAN_ACCENT = RGBColor(56, 189, 248) # #38BDF8
    TEXT_WHITE = RGBColor(248, 250, 252) # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    ALERT_RED = RGBColor(220, 38, 38)    # #DC2626
    ALERT_AMBER = RGBColor(245, 158, 11) # #F59E0B

    def add_blank_slide_with_bg(prs, bg_color):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = bg_color
        bg.line.fill.background()
        return slide

    def add_header(slide, kicker, title):
        kicker_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.35))
        tf_k = kicker_box.text_frame
        tf_k.word_wrap = True
        p_k = tf_k.paragraphs[0]
        p_k.text = kicker.upper()
        p_k.font.size = Pt(11)
        p_k.font.bold = True
        p_k.font.color.rgb = EMERALD
        p_k.font.name = "Arial"

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.75))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.font.name = "Arial"

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(30, 58, 104)
        line.line.fill.background()

    def add_card(slide, left, top, width, height, title, points, border_color=EMERALD, bg_color=PANEL_BG):
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = border_color
        p_title.font.name = "Arial"

        for pt in points:
            p = tf.add_paragraph()
            p.text = "• " + pt
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_WHITE
            p.font.name = "Arial"
            p.space_before = Pt(5)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Green & Blue Theme)
    # -------------------------------------------------------------
    slide1 = add_blank_slide_with_bg(prs, NAVY_BG)
    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = EMERALD
    top_bar.line.fill.background()

    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(3.0))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "AI-POWERED FINANCIAL TRANSACTION DEFENSE"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = EMERALD

    p_main = tf1.add_paragraph()
    p_main.text = "FraudGuard AI"
    p_main.font.size = Pt(48)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_WHITE
    p_main.space_before = Pt(8)

    p_tagline = tf1.add_paragraph()
    p_tagline.text = "Detect. Analyze. Protect."
    p_tagline.font.size = Pt(22)
    p_tagline.font.bold = True
    p_tagline.font.color.rgb = CYAN_ACCENT
    p_tagline.space_before = Pt(4)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Multi-Factor Anomaly Scoring &bull; 📱 User Notification &bull; 📞 User Call Verification &bull; 🔒 Auto-Hold &bull; 👨‍💼 Admin Review &bull; Block/Unhold Defense"
    p_desc.font.size = Pt(14)
    p_desc.font.color.rgb = TEXT_MUTED
    p_desc.space_before = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 2: Core Defense Lifecycle Flowchart (Exact Prompt Architecture)
    # -------------------------------------------------------------
    slide2 = add_blank_slide_with_bg(prs, NAVY_BG)
    add_header(slide2, "CORE WORKFLOW ARCHITECTURE", "End-to-End Incident Escalation & Decision Tree")

    flow_text = (
        "TRANSACTION INGESTION\n"
        "         ↓\n"
        "AI DETECTS HIGH RISK (Scikit-Learn ML + Rules)\n"
        "         ↓\n"
        "📱 USER NOTIFICATION (SMS / WhatsApp)\n"
        "         ↓\n"
        "📞 CALL USER (Attempt #1)\n"
        "         ↓\n"
        "Call attended? ──→ [YES] ──→ Verify (Safe / Block)\n"
        "         ↓ [NO]\n"
        "   RETRY CALL (Attempt #2)\n"
        "         ↓\n"
        "Call attended? ──→ [YES] ──→ Verify (Safe / Block)\n"
        "         ↓ [NO]\n"
        "🔒 AUTOMATED ACCOUNT HOLD\n"
        "         ↓\n"
        "👨‍💼 ADMIN REVIEW QUEUE\n"
        "   ┌──────────────┴──────────────┐\n"
        "   ↓                             ↓\n"
        " [SAFE]                       [FRAUD]\n"
        "   ↓                             ↓\n"
        "UNHOLD ACCOUNT               BLOCK ACCOUNT"
    )

    card_f = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    card_f.fill.solid()
    card_f.fill.fore_color.rgb = PANEL_BG
    card_f.line.color.rgb = EMERALD
    card_f.line.width = Pt(2)

    tb_f = slide2.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11.0), Inches(4.5))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True
    p_flow = tf_f.paragraphs[0]
    p_flow.text = flow_text
    p_flow.font.size = Pt(12.5)
    p_flow.font.name = "Consolas"
    p_flow.font.color.rgb = CYAN_ACCENT

    # -------------------------------------------------------------
    # SLIDE 3: Problem Statement
    # -------------------------------------------------------------
    slide3 = add_blank_slide_with_bg(prs, NAVY_BG)
    add_header(slide3, "THE PROBLEM", "Challenges in Modern Fraud Prevention")

    add_card(slide3, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), 
             "Speed of Theft", 
             [
                 "Fraudsters drain accounts in rapid bursts before users notice.",
                 "Traditional batch audits detect crime hours after funds leave.",
                 "Zero-day techniques bypass static rule thresholds."
             ], ALERT_RED)

    add_card(slide3, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8), 
             "Verification Gap", 
             [
                 "Banks struggle to confirm transactions with cardholders instantly.",
                 "Single missed phone calls result in either false declines or stolen balances.",
                 "Lack of structured retry & automated hold protocols."
             ], ALERT_AMBER)

    add_card(slide3, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8), 
             "Investigator Fatigue", 
             [
                 "Opaque AI alert scores with no explainable evidence.",
                 "No unified admin review queue to triage held accounts.",
                 "Manual processes slow down emergency card freezing."
             ], CYAN_ACCENT)

    # -------------------------------------------------------------
    # SLIDE 4: Proposed Solution
    # -------------------------------------------------------------
    slide4 = add_blank_slide_with_bg(prs, NAVY_BG)
    add_header(slide4, "THE SOLUTION", "FraudGuard AI Defense Innovation")

    add_card(slide4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "1. Dynamic Risk & Isolation Forest",
             [
                 "Scikit-Learn ML model evaluating multi-dimensional outlier depth.",
                 "Weighted scoring: 60% heuristic rules + 40% ML anomaly score."
             ], EMERALD)

    add_card(slide4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.3),
             "2. 📱 SMS & 📞 2-Stage Voice Call",
             [
                 "Dispatches immediate user SMS notification upon High Risk.",
                 "Automated outbound voice call with retry attempt logic."
             ], ROYAL_BLUE)

    add_card(slide4, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.4),
             "3. 🔒 Automated Fail-Safe Hold",
             [
                 "Locks account transfers if 2 calls go unanswered to prevent drain.",
                 "Maintains security without permanently shutting down accounts."
             ], ALERT_AMBER)

    add_card(slide4, Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.4),
             "4. 👨‍💼 Admin Adjudication Queue",
             [
                 "Security analyst reviews evidence & unholds or blocks account.",
                 "Explainable AI factors (01, 02, 03) guide rapid decision-making."
             ], ALERT_RED)

    # -------------------------------------------------------------
    # SLIDE 5: Machine Learning & Scoring Formulation
    # -------------------------------------------------------------
    slide5 = add_blank_slide_with_bg(prs, NAVY_BG)
    add_header(slide5, "TECHNICAL METHODOLOGY", "Scoring Formulation & Machine Learning")

    add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Heuristic Penalty Allocation (Max: 100)",
             [
                 "Amount Multiplier (+25): Amount > 3× account baseline average.",
                 "Hardware Signature (+20): Unregistered device / browser fingerprint.",
                 "Geographic Shift (+20): Foreign or unverified IP corridor.",
                 "Off-Hours Timestamp (+15): 11:00 PM – 5:30 AM execution.",
                 "Burst Velocity (+20): ≥ 5 transactions in 10 minutes."
             ], EMERALD)

    add_card(slide5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8),
             "Isolation Forest & Classification",
             [
                 "Model: Scikit-Learn IsolationForest(120 Estimators, Contamination=0.08).",
                 "Formula: Final_Risk = round(0.60 × Rule + 0.40 × ML_Score).",
                 "Tiers:",
                 "  • LOW (0–30): Verified Safe Standard Activity",
                 "  • MEDIUM (31–60): Secondary MFA Review",
                 "  • HIGH (61–80): 📱 SMS + 📞 Call Sequence Triggered",
                 "  • CRITICAL (81–100): Immediate Auto-Hold & Review"
             ], ROYAL_BLUE)

    # -------------------------------------------------------------
    # SLIDE 6: UI Design & Network Graph
    # -------------------------------------------------------------
    slide6 = add_blank_slide_with_bg(prs, NAVY_BG)
    add_header(slide6, "UI & NETWORK INTELLIGENCE", "Green & Blue Visual Identity + Mule Ring Graph")

    add_card(slide6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Green & Blue Financial Identity",
             [
                 "Strict Rectangular UI (4px-8px radius): Zero pills, zero capsules.",
                 "Emerald Green & Tech Blue Palette: Professional depth.",
                 "Real-Time Responsive: Adapts cleanly on desktop, tablet, and mobile."
             ], EMERALD)

    add_card(slide6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8),
             "Transaction Relationship Canvas Graph",
             [
                 "HTML5 Canvas Engine: Visualizes account transfer flows.",
                 "Mule Ring Detection: Highlights suspicious coordinated loops.",
                 "Click-to-Inspect: Drill directly into connected account baselines."
             ], CYAN_ACCENT)

    # -------------------------------------------------------------
    # SLIDE 7: Engineering Stack
    # -------------------------------------------------------------
    slide7 = add_blank_slide_with_bg(prs, NAVY_BG)
    add_header(slide7, "ENGINEERING ARCHITECTURE", "Production Full-Stack Technology Stack")

    add_card(slide7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Backend & APIs",
             [
                 "Python 3.13 & Flask 3.1",
                 "SQLite3 Transaction Ledger",
                 "15+ RESTful JSON Endpoints",
                 "Session Auth & Role Control"
             ], ROYAL_BLUE)

    add_card(slide7, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "ML & Vector Core",
             [
                 "Scikit-Learn Isolation Forest",
                 "Pandas & NumPy Pipelines",
                 "Joblib Serialized Models",
                 "7-Feature Anomaly Space"
             ], EMERALD)

    add_card(slide7, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Frontend & Visuals",
             [
                 "Handcrafted Green/Blue CSS",
                 "Chart.js 4.4 Analytics",
                 "HTML5 Canvas Graph Engine",
                 "Live Multi-Stage State Machine"
             ], CYAN_ACCENT)

    # -------------------------------------------------------------
    # SLIDE 8: Hackathon Demo Walkthrough
    # -------------------------------------------------------------
    slide8 = add_blank_slide_with_bg(prs, NAVY_BG)
    add_header(slide8, "JURY DEMONSTRATION", "Live Evaluation Walkthrough")

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Live Decision Tree Demonstration",
             [
                 "1. Click [ Simulate Transaction ] &rarr; AI detects High Risk.",
                 "2. 📱 User Notification SMS is automatically logged & sent.",
                 "3. Click [ Verify Call ] &rarr; Attempt #1 placed to cardholder.",
                 "4. Simulate [ NO - Missed ] &rarr; Auto-triggers Retry Call #2.",
                 "5. Simulate [ NO - Missed Again ] &rarr; 🔒 AUTOMATIC ACCOUNT HOLD.",
                 "6. 👨‍💼 Admin Review Queue: Click [ Unhold ] or [ Block Account ]."
             ], EMERALD)

    add_card(slide8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8),
             "Batch Ingestion & Reporting",
             [
                 "Drag & drop CSV files for instant multi-record scoring.",
                 "Download complete compliance CSV audit trail reports.",
                 "Demonstrates zero fund drainage and investigator efficiency."
             ], ROYAL_BLUE)

    # -------------------------------------------------------------
    # SLIDE 9: Conclusion
    # -------------------------------------------------------------
    slide9 = add_blank_slide_with_bg(prs, NAVY_BG)
    
    t_box = slide9.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.3), Inches(4.5))
    tf9 = t_box.text_frame
    tf9.word_wrap = True

    p_c1 = tf9.paragraphs[0]
    p_c1.text = "FraudGuard AI"
    p_c1.font.size = Pt(42)
    p_c1.font.bold = True
    p_c1.font.color.rgb = EMERALD

    p_c2 = tf9.add_paragraph()
    p_c2.text = "Detect. Analyze. Protect."
    p_c2.font.size = Pt(22)
    p_c2.font.bold = True
    p_c2.font.color.rgb = CYAN_ACCENT
    p_c2.space_before = Pt(6)

    p_c3 = tf9.add_paragraph()
    p_c3.text = "A complete, end-to-end transaction security ecosystem featuring automated call retries, fail-safe account holds, and human-in-the-loop admin review."
    p_c3.font.size = Pt(16)
    p_c3.font.color.rgb = TEXT_WHITE
    p_c3.space_before = Pt(16)

    p_c4 = tf9.add_paragraph()
    p_c4.text = "Thank you! Ready for Live Jury Evaluation."
    p_c4.font.size = Pt(20)
    p_c4.font.bold = True
    p_c4.font.color.rgb = EMERALD
    p_c4.space_before = Pt(24)

    output_path = os.path.join(os.path.dirname(__file__), 'FraudGuard_AI_Presentation.pptx')
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

    desktop_path = r'C:\Users\Dell\Desktop\fraudguard-ai\FraudGuard_AI_Presentation.pptx'
    try:
        prs.save(desktop_path)
        print(f"Presentation saved to Desktop: {desktop_path}")
    except Exception as e:
        print(f"Could not save directly to desktop: {e}")

if __name__ == '__main__':
    create_presentation()
